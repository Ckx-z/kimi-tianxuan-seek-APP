"""实验进展 progress（v2.2.0）测试：三选一互斥 + 筛选 + 旧数据兼容 + draft/final 不受影响。

覆盖：
- 旧记录 json 缺 progress 字段 → 读出补默认 running（读取兼容不落盘改写）；
- 三个取值均可创建 / 更新并持久化（重新读回一致，落盘文件里也有该字段）；
- 非法值被拒绝：store 层 ValueError（中文提示列出允许值），API 层 400；
- 列表按 progress 筛选得到正确子集（running / finished / xrd_done 各测一次）；
- 现有 draft/final 语义未被破坏：草稿宽松校验仍生效、final 完整校验不变，
  且 progress 非法时不论草稿还是正式都报错；
- Word 导出同时保留「状态」行并新增「实验进展」行。

数据隔离：所有写操作经 monkeypatch 把收藏/记录/附件目录指到 pytest tmp_path
（--basetemp 落在仓库 .tmp_xtb/ 内），并用 tmp 数据目录，不碰仓库 data/
或真实用户数据目录。
"""

from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from favorites import store as fav_store  # noqa: E402
from records import store as rec_store  # noqa: E402
from api.main import app  # noqa: E402

client = TestClient(app)

TP = "O=CC1=C(C=O)C(=O)C(C=O)=C1O"
PA = "Nc1ccc(N)cc1"

# 三种进展取值（与后端常量同源，避免写死漂移）
RUNNING, FINISHED, XRD_DONE = rec_store.PROGRESS_VALUES


@pytest.fixture
def rec_dirs(tmp_path, monkeypatch):
    """收藏 / 记录 / 附件目录全部指到 tmp_path，杜绝污染真实数据。"""
    fav_d = tmp_path / "favorites"
    rec_d = tmp_path / "records"
    att_d = tmp_path / "attachments"
    monkeypatch.setattr(fav_store, "FAVORITES_DIR", fav_d)
    monkeypatch.setattr(rec_store, "RECORDS_DIR", rec_d)
    monkeypatch.setattr(rec_store, "ATTACHMENTS_DIR", att_d)
    return fav_d, rec_d, att_d


def _write_legacy_record(rec_d: Path, rec_id: str = "rec_20250101_001",
                         experiment_no: str = "A1") -> Path:
    """写一条「老记录」：完全不含 progress 字段（模拟 v2.2.0 之前的存量数据）。"""
    rec_d.mkdir(parents=True, exist_ok=True)
    legacy = {
        "schema_version": "1.0",
        "record_type": "experiment_record",
        "record_id": rec_id,
        "experiment_no": experiment_no,
        "status": "final",
        "favorite_id": None,
        "aldehyde": {"smiles": TP, "cas": "", "name": ""},
        "amine": {"smiles": PA, "cas": "", "name": ""},
        "prediction_id": None,
        "prediction_snapshot": None,
        "conditions": {},
        "outcome": "film",
        "failure_class": None,
        "strength": "",
        "notes": f"实验编号：{experiment_no}",
        "process_notes": "",
        "timeline": [],
        "attachments": [],
        "operator": "",
        "date": "2025-01-01",
        "minimax_plan_no": None,
    }
    path = rec_d / f"{rec_id}.json"
    path.write_text(json.dumps(legacy, ensure_ascii=False, indent=1) + "\n",
                    encoding="utf-8")
    return path


def _new_favorite() -> str:
    fav = fav_store.add_favorite(TP, PA)
    return fav["id"]


def _on_disk(rec_d: Path, rec_id: str) -> dict:
    return json.loads((rec_d / f"{rec_id}.json").read_text(encoding="utf-8"))


# ---------------------------------------------------------------------------
# 1. 常量与旧数据兼容
# ---------------------------------------------------------------------------


class TestProgressConstantsAndLegacyCompat:
    def test_values_and_chinese_labels(self):
        assert rec_store.PROGRESS_VALUES == ("running", "finished", "xrd_done")
        assert rec_store.DEFAULT_PROGRESS == "running"
        assert rec_store.PROGRESS_LABELS == {
            "running": "实验进行中",
            "finished": "实验结束",
            "xrd_done": "XRD 已完成",
        }

    def test_legacy_record_missing_field_reads_running(self, rec_dirs):
        _, rec_d, _ = rec_dirs
        path = _write_legacy_record(rec_d)
        # 前置事实：老记录文件里确实没有 progress 字段
        assert "progress" not in json.loads(path.read_text(encoding="utf-8"))

        got = rec_store.get_record("rec_20250101_001")
        assert got is not None and got["progress"] == RUNNING
        # 列表路径同样补默认值
        assert rec_store.list_records()[0]["progress"] == RUNNING
        # 读取兼容只作用于返回值，不改写用户文件
        assert "progress" not in json.loads(path.read_text(encoding="utf-8"))

    def test_legacy_record_empty_or_broken_value_reads_running(self, rec_dirs):
        _, rec_d, _ = rec_dirs
        _write_legacy_record(rec_d, "rec_20250101_002", "A2")
        raw = _on_disk(rec_d, "rec_20250101_002")
        raw["progress"] = ""  # 空串（历史写入异常）
        (rec_d / "rec_20250101_002.json").write_text(
            json.dumps(raw, ensure_ascii=False), encoding="utf-8")
        assert rec_store.get_record("rec_20250101_002")["progress"] == RUNNING


# ---------------------------------------------------------------------------
# 2. 三取值创建 / 更新 + 持久化
# ---------------------------------------------------------------------------


class TestProgressCreateUpdatePersist:
    @pytest.mark.parametrize("value", [RUNNING, FINISHED, XRD_DONE])
    def test_create_with_each_value_persists(self, rec_dirs, value):
        _, rec_d, _ = rec_dirs
        rec = rec_store.create_record(
            favorite_id=_new_favorite(), experiment_no="A1", outcome="film",
            progress=value)
        assert rec["progress"] == value
        rid = rec["record_id"]
        assert rec_store.get_record(rid)["progress"] == value      # 重新读回一致
        assert _on_disk(rec_d, rid)["progress"] == value           # 真的落盘了
        assert rec_store.list_records(progress=value)[0]["record_id"] == rid

    def test_create_without_progress_defaults_running(self, rec_dirs):
        """向后兼容：不传 progress 的旧调用照常工作，落盘为 running。"""
        _, rec_d, _ = rec_dirs
        rec = rec_store.create_record(
            favorite_id=_new_favorite(), experiment_no="A1", outcome="film")
        assert rec["progress"] == RUNNING
        assert _on_disk(rec_d, rec["record_id"])["progress"] == RUNNING

    @pytest.mark.parametrize("value", [RUNNING, FINISHED, XRD_DONE])
    def test_update_with_each_value_persists(self, rec_dirs, value):
        _, rec_d, _ = rec_dirs
        rec = rec_store.create_record(
            favorite_id=_new_favorite(), experiment_no="A1", outcome="film")
        rid = rec["record_id"]
        updated = rec_store.update_record(rid, {"progress": value})
        assert updated["progress"] == value
        assert rec_store.get_record(rid)["progress"] == value      # 重新读回一致
        assert _on_disk(rec_d, rid)["progress"] == value

    def test_update_other_fields_keeps_progress(self, rec_dirs):
        _, rec_d, _ = rec_dirs
        rec = rec_store.create_record(
            favorite_id=_new_favorite(), experiment_no="A1", outcome="film",
            progress=XRD_DONE)
        rid = rec["record_id"]
        updated = rec_store.update_record(rid, {"strength": "柔韧"})
        assert updated["progress"] == XRD_DONE
        assert rec_store.get_record(rid)["progress"] == XRD_DONE

    def test_update_legacy_record_adds_progress(self, rec_dirs):
        """老记录（无该字段）经更新后可落盘新进展。"""
        _, rec_d, _ = rec_dirs
        _write_legacy_record(rec_d)
        updated = rec_store.update_record(
            "rec_20250101_001", {"progress": FINISHED})
        assert updated["progress"] == FINISHED
        assert _on_disk(rec_d, "rec_20250101_001")["progress"] == FINISHED
        assert rec_store.get_record("rec_20250101_001")["progress"] == FINISHED


# ---------------------------------------------------------------------------
# 3. 非法值被拒绝
# ---------------------------------------------------------------------------


class TestProgressInvalidRejected:
    def test_create_invalid_raises_valueerror_with_allowed_values(self, rec_dirs):
        with pytest.raises(ValueError) as exc:
            rec_store.create_record(
                favorite_id=_new_favorite(), experiment_no="A1",
                outcome="film", progress="done")
        msg = str(exc.value)
        assert "progress" in msg and "实验进展" in msg
        for value in rec_store.PROGRESS_VALUES:  # 提示里列出全部允许值
            assert value in msg

    def test_update_invalid_raises_and_keeps_original(self, rec_dirs):
        _, rec_d, _ = rec_dirs
        rec = rec_store.create_record(
            favorite_id=_new_favorite(), experiment_no="A1", outcome="film",
            progress=FINISHED)
        rid = rec["record_id"]
        with pytest.raises(ValueError):
            rec_store.update_record(rid, {"progress": "bogus"})
        assert rec_store.get_record(rid)["progress"] == FINISHED  # 未被写坏
        assert _on_disk(rec_d, rid)["progress"] == FINISHED

    def test_list_filter_invalid_raises(self, rec_dirs):
        with pytest.raises(ValueError):
            rec_store.list_records(progress="bogus")


# ---------------------------------------------------------------------------
# 4. 列表按 progress 筛选（三种各测一次）
# ---------------------------------------------------------------------------


class TestProgressListFilter:
    @pytest.fixture
    def mixed_records(self, rec_dirs):
        """三条记录各一个进展 + 一条缺字段老记录（应算 running）。"""
        _, rec_d, _ = rec_dirs
        fid = _new_favorite()
        made = {}
        for value, no in ((RUNNING, "R1"), (FINISHED, "F1"), (XRD_DONE, "X1")):
            rec = rec_store.create_record(
                favorite_id=fid, experiment_no=no, outcome="film",
                progress=value)
            made[value] = rec["record_id"]
        _write_legacy_record(rec_d, "rec_20250101_009", "L1")
        made["legacy"] = "rec_20250101_009"
        return made

    def test_filter_running(self, mixed_records):
        got = rec_store.list_records(progress=RUNNING)
        ids = {r["record_id"] for r in got}
        assert ids == {mixed_records[RUNNING], mixed_records["legacy"]}
        assert all(r["progress"] == RUNNING for r in got)

    def test_filter_finished(self, mixed_records):
        got = rec_store.list_records(progress=FINISHED)
        assert {r["record_id"] for r in got} == {mixed_records[FINISHED]}
        assert all(r["progress"] == FINISHED for r in got)

    def test_filter_xrd_done(self, mixed_records):
        got = rec_store.list_records(progress=XRD_DONE)
        assert {r["record_id"] for r in got} == {mixed_records[XRD_DONE]}
        assert all(r["progress"] == XRD_DONE for r in got)

    def test_no_filter_returns_all(self, mixed_records):
        assert len(rec_store.list_records()) == 4
        assert len(rec_store.list_records(progress=None)) == 4
        assert len(rec_store.list_records(progress="")) == 4  # 空串视为不筛选

    def test_filter_combined_with_favorite_id(self, mixed_records, rec_dirs):
        fav_id = rec_store.get_record(mixed_records[FINISHED])["favorite_id"]
        got = rec_store.list_records(favorite_id=fav_id, progress=XRD_DONE)
        assert {r["record_id"] for r in got} == {mixed_records[XRD_DONE]}


# ---------------------------------------------------------------------------
# 5. API 层：筛选 / 创建 / 更新 / 非法值 400 / 向后兼容
# ---------------------------------------------------------------------------


class TestProgressApi:
    def _create_via_api(self, fid: str, no: str, progress: str | None = None) -> dict:
        body: dict = {"favorite_id": fid, "experiment_no": no, "outcome": "film"}
        if progress is not None:
            body["progress"] = progress
        r = client.post("/api/records", json=body)
        assert r.status_code == 201, r.text
        return r.json()

    def test_create_returns_progress_and_defaults_running(self, rec_dirs):
        fid = client.post("/api/favorites", json={
            "aldehyde_smiles": TP, "amine_smiles": PA}).json()["id"]
        # 不传 progress（旧调用）→ 响应/详情/列表均为 running
        rec = self._create_via_api(fid, "A1")
        assert rec["progress"] == RUNNING
        rid = rec["record_id"]
        assert client.get(f"/api/records/{rid}").json()["progress"] == RUNNING
        listed = client.get("/api/records").json()["records"]
        assert listed[0]["progress"] == RUNNING
        # 显式传 finished
        rec2 = self._create_via_api(fid, "A2", FINISHED)
        assert rec2["progress"] == FINISHED

    def test_update_returns_progress(self, rec_dirs):
        fid = client.post("/api/favorites", json={
            "aldehyde_smiles": TP, "amine_smiles": PA}).json()["id"]
        rid = self._create_via_api(fid, "A1")["record_id"]
        r = client.put(f"/api/records/{rid}", json={"progress": XRD_DONE})
        assert r.status_code == 200, r.text
        assert r.json()["progress"] == XRD_DONE
        assert client.get(f"/api/records/{rid}").json()["progress"] == XRD_DONE
        # 不含 progress 的更新请求不改变进展（向后兼容）
        r = client.put(f"/api/records/{rid}", json={"notes": "改备注"})
        assert r.status_code == 200
        assert r.json()["progress"] == XRD_DONE

    def test_list_filter_by_query_param(self, rec_dirs):
        fid = client.post("/api/favorites", json={
            "aldehyde_smiles": TP, "amine_smiles": PA}).json()["id"]
        self._create_via_api(fid, "R1", RUNNING)
        self._create_via_api(fid, "F1", FINISHED)
        self._create_via_api(fid, "X1", XRD_DONE)
        for value, no in ((RUNNING, "R1"), (FINISHED, "F1"), (XRD_DONE, "X1")):
            r = client.get(f"/api/records?progress={value}")
            assert r.status_code == 200, r.text
            got = r.json()["records"]
            assert [x["experiment_no"] for x in got] == [no]
        # 无参数 = 全部
        assert len(client.get("/api/records").json()["records"]) == 3

    def test_list_invalid_progress_returns_400(self, rec_dirs):
        r = client.get("/api/records?progress=bogus")
        assert r.status_code == 400
        detail = r.json()["detail"]
        assert "progress" in detail and "实验进展" in detail
        for value in rec_store.PROGRESS_VALUES:
            assert value in detail

    def test_create_invalid_progress_returns_400(self, rec_dirs):
        fid = client.post("/api/favorites", json={
            "aldehyde_smiles": TP, "amine_smiles": PA}).json()["id"]
        r = client.post("/api/records", json={
            "favorite_id": fid, "experiment_no": "A1", "outcome": "film",
            "progress": "done"})
        assert r.status_code == 400
        assert "progress" in r.json()["detail"]

    def test_update_invalid_progress_returns_400(self, rec_dirs):
        fid = client.post("/api/favorites", json={
            "aldehyde_smiles": TP, "amine_smiles": PA}).json()["id"]
        rid = self._create_via_api(fid, "A1", FINISHED)["record_id"]
        r = client.put(f"/api/records/{rid}", json={"progress": "done"})
        assert r.status_code == 400
        assert client.get(f"/api/records/{rid}").json()["progress"] == FINISHED


# ---------------------------------------------------------------------------
# 6. draft/final 既有语义未被破坏
# ---------------------------------------------------------------------------


class TestStatusBehaviorUnchanged:
    def test_draft_lenient_validation_still_works(self, rec_dirs):
        _, rec_d, _ = rec_dirs
        rec = rec_store.create_record(
            favorite_id=_new_favorite(), status="draft", progress=XRD_DONE)
        assert rec["status"] == "draft"
        assert rec["progress"] == XRD_DONE
        assert rec["experiment_no"] == ""      # 草稿编号可空
        assert rec["outcome"] == ""            # 草稿结果可空
        assert not rec["notes"].startswith("实验编号：")
        # 更新路径：草稿仍宽松
        updated = rec_store.update_record(rec["record_id"], {"strength": "脆"})
        assert updated["status"] == "draft"
        assert updated["progress"] == XRD_DONE

    def test_final_validation_still_strict(self, rec_dirs):
        fid = _new_favorite()
        with pytest.raises(ValueError):
            rec_store.create_record(favorite_id=fid, status="final")
        with pytest.raises(ValueError):
            rec_store.create_record(
                favorite_id=fid, status="final", experiment_no="A1")
        # status 非法值照旧报错（与 progress 校验互不影响）
        with pytest.raises(ValueError):
            rec_store.create_record(
                favorite_id=fid, status="wip", experiment_no="A1",
                outcome="film", progress=RUNNING)
        # progress 非法时，草稿与正式都报错
        with pytest.raises(ValueError):
            rec_store.create_record(
                favorite_id=fid, status="draft", progress="done")
        with pytest.raises(ValueError):
            rec_store.create_record(
                favorite_id=fid, status="final", experiment_no="A1",
                outcome="film", progress="done")

    def test_api_draft_to_final_flow_unchanged(self, rec_dirs):
        fid = client.post("/api/favorites", json={
            "aldehyde_smiles": TP, "amine_smiles": PA}).json()["id"]
        r = client.post("/api/records", json={
            "favorite_id": fid, "status": "draft", "progress": FINISHED})
        assert r.status_code == 201, r.text
        rec = r.json()
        rid = rec["record_id"]
        assert rec["status"] == "draft" and rec["progress"] == FINISHED
        # 缺 outcome 转正式 → 400（草稿宽松语义不变）
        assert client.put(f"/api/records/{rid}", json={"status": "final"}).status_code == 400
        # 补齐转正式 → 200，notes 加编号前缀，progress 保持
        r = client.put(f"/api/records/{rid}", json={
            "status": "final", "experiment_no": "A5", "outcome": "film"})
        assert r.status_code == 200, r.text
        final = r.json()
        assert final["status"] == "final"
        assert final["notes"].startswith("实验编号：A5")
        assert final["progress"] == FINISHED


# ---------------------------------------------------------------------------
# 7. Word 导出：保留「状态」行 + 新增「实验进展」行
# ---------------------------------------------------------------------------


class TestProgressInDocxExport:
    @pytest.fixture
    def no_llm(self, monkeypatch):
        """性质卡生成置空（未配置 LLM 的降级路径），保持测试离线、快速。"""
        import recommend.monomer_props as monomer_props
        monkeypatch.setattr(monomer_props, "get_monomer_properties",
                            lambda *a, **k: None, raising=False)

    def _docx_text(self, rec: dict) -> str:
        from docx import Document

        from records import export_docx
        data = export_docx.build_record_docx(rec, version="test")
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)

    def test_docx_has_progress_row_and_keeps_status_row(self, rec_dirs, no_llm):
        rec = rec_store.create_record(
            favorite_id=_new_favorite(), experiment_no="A5", outcome="film",
            progress=XRD_DONE)
        text = self._docx_text(rec)
        assert "实验进展" in text
        assert "XRD 已完成" in text
        # 原有「状态」行（draft/final）保持不变
        assert "状态" in text and "正式" in text
